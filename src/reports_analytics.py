import os
import sqlite3
import csv
from datetime import datetime
from decimal import Decimal
from typing import Dict, List, Tuple

from openpyxl import Workbook

DB_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "invoices.db")


Summary = Dict[str, float]
InvoiceRow = Dict[str, str]


def _fetch_quantity_map(conn: sqlite3.Connection, invoice_ids: List[int]) -> Dict[int, Decimal]:
    if not invoice_ids:
        return {}
    placeholders = ",".join("?" for _ in invoice_ids)
    query = f"""
        SELECT invoice_id, COALESCE(SUM(qty), 0) AS qty_sum
          FROM invoice_items
         WHERE invoice_id IN ({placeholders})
         GROUP BY invoice_id
    """
    rows = conn.execute(query, invoice_ids).fetchall()
    return {row["invoice_id"]: Decimal(str(row["qty_sum"])) for row in rows}


def fetch_summary(start_date: str, end_date: str) -> Tuple[Summary, List[InvoiceRow]]:
    """Return totals and invoice detail for a given period."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    query = """
        SELECT i.id,
               i.invoice_no,
               i.date,
               i.subtotal,
               COALESCE(i.sales_tax, 0)     AS sales_tax_total,
               COALESCE(i.advance_tax, 0)   AS advance_tax_total
          FROM invoices i
         WHERE i.date BETWEEN ? AND ?
         ORDER BY i.date ASC, i.invoice_no ASC
    """
    try:
        rows = [dict(r) for r in conn.execute(query, (start_date, end_date))]
        qty_map = _fetch_quantity_map(conn, [row["id"] for row in rows])
    finally:
        conn.close()

    total_sales = Decimal("0")
    total_sales_tax = Decimal("0")
    total_adv_tax = Decimal("0")
    total_qty = Decimal("0")

    for inv in rows:
        inv_id = inv["id"]
        inv_qty = qty_map.get(inv_id, Decimal("0"))
        inv["total_quantity"] = float(inv_qty)

        total_sales += Decimal(str(inv.get("subtotal", 0)))
        total_sales_tax += Decimal(str(inv.get("sales_tax_total", 0)))
        total_adv_tax += Decimal(str(inv.get("advance_tax_total", 0)))
        total_qty += inv_qty

    summary: Summary = {
        "period_start": start_date,
        "period_end": end_date,
        "total_sales": float(total_sales),
        "total_sales_tax": float(total_sales_tax),
        "total_advance_tax": float(total_adv_tax),
        "total_quantity": float(total_qty),
        "invoice_count": len(rows),
    }
    return summary, rows


def fetch_top_products(start_date: str, end_date: str, limit: int = 5) -> List[Dict]:
    """Return top selling products by revenue for the period."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    
    # We group by product_id and join products to get the current name.
    # If product_id is null, we skip it for now or could group by description.
    query = """
        SELECT p.name as product_name,
               SUM(ii.qty) as total_qty,
               SUM(ii.total_amount) as total_revenue
        FROM invoice_items ii
        JOIN invoices i ON i.id = ii.invoice_id
        LEFT JOIN products p ON p.id = ii.product_id
        WHERE i.date BETWEEN ? AND ?
        GROUP BY ii.product_id
        HAVING product_name IS NOT NULL
        ORDER BY total_revenue DESC
        LIMIT ?
    """
    try:
        rows = conn.execute(query, (start_date, end_date, limit)).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def fetch_top_customers(start_date: str, end_date: str, limit: int = 5) -> List[Dict]:
    """Return top spending customers for the period."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    
    query = """
        SELECT c.name as customer_name,
               COUNT(i.id) as invoice_count,
               SUM(i.total_amount) as total_spent
        FROM invoices i
        JOIN customers c ON c.id = i.customer_id
        WHERE i.date BETWEEN ? AND ?
        GROUP BY c.id
        ORDER BY total_spent DESC
        LIMIT ?
    """
    try:
        rows = conn.execute(query, (start_date, end_date, limit)).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


# --------------------------
# CSV export
# --------------------------
def export_summary_to_csv(summary, rows, output_path):
    """Write period summary and invoice detail to CSV."""
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["INVOICE REPORT"])
        writer.writerow([f"Period: {summary['period_start']} to {summary['period_end']}"])
        writer.writerow([])
        writer.writerow(["Invoice No", "Date", "Subtotal", "Sales Tax", "Advance Tax"])
        for inv in rows:
            writer.writerow([
                inv["invoice_no"],
                inv["date"],
                inv["subtotal"],
                inv["sales_tax_total"],
                inv["advance_tax_total"],
            ])
        writer.writerow([])
        writer.writerow(["TOTAL SALES", summary["total_sales"]])
        writer.writerow(["TOTAL SALES TAX", summary["total_sales_tax"]])
        writer.writerow(["TOTAL ADVANCE TAX", summary["total_advance_tax"]])
        writer.writerow(["TOTAL QUANTITY (pcs)", summary["total_quantity"]])
        writer.writerow(["NUMBER OF INVOICES", summary["invoice_count"]])
    return output_path


# --------------------------
# Excel export
# --------------------------
def export_summary_to_excel(summary, rows, output_path):
    """Write same report to .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Summary"

    ws.append(["INVOICE REPORT"])
    ws.append([f"Period: {summary['period_start']} to {summary['period_end']}"])
    ws.append([])

    ws.append(["Invoice No", "Date", "Subtotal", "Sales Tax", "Advance Tax"])
    for inv in rows:
        ws.append([
            inv["invoice_no"],
            inv["date"],
            inv["subtotal"],
            inv["sales_tax_total"],
            inv["advance_tax_total"],
        ])
    ws.append([])

    ws.append(["TOTAL SALES", summary["total_sales"]])
    ws.append(["TOTAL SALES TAX", summary["total_sales_tax"]])
    ws.append(["TOTAL ADVANCE TAX", summary["total_advance_tax"]])
    ws.append(["TOTAL QUANTITY (pcs)", summary["total_quantity"]])
    ws.append(["NUMBER OF INVOICES", summary["invoice_count"]])

    wb.save(output_path)
    return output_path


# --------------------------
# Quick manual test
# --------------------------
if __name__ == "__main__":
    start, end = "2025-10-01", "2025-10-31"
    summary, rows = fetch_summary(start, end)
    print(summary)
    export_summary_to_csv(summary, rows, "monthly_report.csv")
    export_summary_to_excel(summary, rows, "monthly_report.xlsx")
    print("Reports exported successfully.")